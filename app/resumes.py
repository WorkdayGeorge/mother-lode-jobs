"""Reading resumes.

Only the extracted text is kept (for matching). The uploaded file itself is
never stored, which keeps the amount of personal data we hold small.
"""
from __future__ import annotations

from io import BytesIO

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 30000


class ResumeError(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    if len(content) > MAX_BYTES:
        raise ResumeError("That file is over 5 MB. Try saving it as a smaller PDF or Word file.")
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:10])
        elif name.endswith(".docx"):
            from docx import Document

            document = Document(BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="replace")
        else:
            raise ResumeError("Upload a PDF, Word (.docx) or plain text file.")
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError("We couldn't read that file. Try saving it as a PDF and uploading again.") from exc
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 40:
        raise ResumeError("We couldn't find any text in that file. If it's a scanned image, upload a Word or text version.")
    return text[:MAX_CHARS]
